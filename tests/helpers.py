"""Test helpers: fake embedder and sample-document builders."""

import hashlib
import os
import re

import numpy as np
from docx import Document

from odf.index import file_indexer


class FakeEmbedder:
    """Deterministic bag-of-words embedder (384-dim, L2-normalised)."""

    DIM = 384

    def __init__(self, *args, **kwargs):
        self.model_name = "fake"

    def _vec(self, text):
        v = np.zeros(self.DIM, dtype=np.float32)
        for tok in re.findall(r"[a-z0-9]+", text.lower()):
            v[int(hashlib.md5(tok.encode()).hexdigest(), 16) % self.DIM] += 1.0
        n = np.linalg.norm(v)
        return v / n if n else v

    def embed_texts(self, texts, batch_size=32):
        if not texts:
            return np.array([])
        return np.array([self._vec(t) for t in texts])

    def embed_text(self, text):
        if not text:
            return np.array([])
        return self._vec(text)


def make_docx(path, paragraphs=(), table=None):
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    doc.save(str(path))
    return str(path)


def make_pdf(path, text):
    """Write a minimal single-page PDF whose text layer contains `text`."""
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    with open(path, "wb") as f:
        f.write(out)
    return str(path)


def bump_mtime(path, seconds=10):
    """Make a file look modified without depending on filesystem timestamp resolution."""
    st = os.stat(path)
    os.utime(path, (st.st_atime, st.st_mtime + seconds))


def index(store, paths):
    """Run the same index pipeline the UI uses (_index_thread)."""
    store.add_documents(file_indexer.process_files(paths, existing_ids=store.get_all_ids()))
