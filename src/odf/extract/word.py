"""Word (.docx) text extraction: paragraphs, then table cells."""

import logging

from docx import Document

from odf.extract.clean import clean_text

log = logging.getLogger(__name__)


def extract(file_path):
    try:
        doc = Document(file_path)
        full_text = [paragraph.text for paragraph in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.append(cell.text)
        return clean_text("\n".join(full_text))
    except Exception as e:
        log.warning("Could not read Word document %s: %s", file_path, e)
        return ""
